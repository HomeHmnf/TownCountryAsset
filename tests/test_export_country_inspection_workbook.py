from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from tools.export_country_inspection_workbook import ISO_HEADER, build_workbook, read_markdown_table


def test_workbook_contains_all_review_sheets_and_keeps_candidates_unapproved(tmp_path: Path) -> None:
    gid = "tc_test"
    registry = {
        "records": [
            {
                "game_id": gid,
                "source": {"current_source_id": "TST", "previous_source_ids": []},
                "status": "proposed",
                "iso": {"alpha2": None, "alpha3": None, "assignment_status": "unverified"},
                "names": {
                    "en": "Test Entity",
                    "fa": None,
                    "aliases_en": ["Example"],
                    "aliases_fa": [],
                    "localization_status": "unreviewed",
                },
                "entity_type": "unreviewed",
                "regions": ["Test Region"],
                "geometry_key": "TST",
            }
        ]
    }
    countries = {
        "count": 1,
        "countries": [
            {
                "game_id": gid,
                "id": "TST",
                "name": "Test Entity",
                "iso3": "TST",
                "regions": ["Test Region"],
                "source_attributes": {"ADMIN": "Test Entity", "CUSTOM_FIELD": "retained"},
            }
        ],
    }
    iso_rows = [
        {
            "game_id": gid,
            "Registry English name": "Test Entity",
            "Candidate alpha-2": "—",
            "Candidate alpha-3": "—",
            "Reference English name": "—",
            "Match type": "no match",
            "Decision": "review required",
        }
    ]
    persian_rows = [
        {
            "game_id": gid,
            "Source ID": "TST",
            "English name": "Test Entity",
            "Current Persian name": "—",
            "Wikidata QID": "—",
            "Candidate Persian name": "—",
            "Candidate status": "missing",
            "Review status": "unreviewed",
        }
    ]
    flags = {
        "records": [
            {
                "game_id": gid,
                "source_flag_code": None,
                "status": "unreviewed",
                "decision_note": None,
                "evidence": [],
            }
        ]
    }
    output = tmp_path / "inspection.xlsx"
    counts = build_workbook(registry, countries, iso_rows, persian_rows, flags, output)

    assert output.exists()
    assert counts["registry_records"] == 1
    assert counts["source_attribute_columns"] == 5
    workbook = load_workbook(output, read_only=False, data_only=True)
    assert workbook.sheetnames == [
        "Read Me",
        "Master Inventory",
        "ISO Review",
        "Entity Review",
        "Persian Review",
        "Flag Review",
        "Data Quality",
        "Source Attributes",
    ]
    inventory = workbook["Master Inventory"]
    headers = [cell.value for cell in inventory[1]]
    values = [cell.value for cell in inventory[2]]
    row = dict(zip(headers, values, strict=True))
    assert row["game_id"] == gid
    assert row["ISO alpha-2 candidate"] == "—"
    assert row["Flag mapping status"] == "unreviewed"
    assert row["Candidate present"] == "yes"
    assert "CUSTOM_FIELD" in [cell.value for cell in workbook["Source Attributes"][1]]
    assert workbook["ISO Review"]["K2"].value is None
    assert workbook["Flag Review"]["K2"].value is None
    assert workbook["Data Quality"].max_row > 1
    workbook.close()



def test_iso_report_parser_accepts_evidence_review_schema_and_normalizes_disposition(tmp_path: Path) -> None:
    report = tmp_path / "iso-assignment-proposals.md"
    report.write_text(
        "\n".join(
            [
                "# ISO proposal review",
                "",
                "| game_id | Registry English name | Candidate alpha-2 | Candidate alpha-3 | Reference English name | Match type | Authoritative evidence status | Proposed disposition |",
                "|---|---|---|---|---|---|---|---|",
                "| tc_example | Example Entity | EX | EXM | Example Entity | exact normalized name | OBP confirmation pending | Keep as unverified candidate |",
                "",
            ]
        ),
        encoding="utf-8",
    )

    rows = read_markdown_table(report, ISO_HEADER, "ISO proposal")

    assert len(rows) == 1
    assert rows[0]["game_id"] == "tc_example"
    assert rows[0]["Candidate alpha-2"] == "EX"
    assert rows[0]["Decision"] == "Keep as unverified candidate"
