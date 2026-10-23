#!/usr/bin/env python3
"""Export a read-only, multi-sheet human inspection workbook for country assets.

The workbook is a review artifact only. It never changes the identity registry,
approves ISO assignments, writes localization candidates back, or approves flags.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REGISTRY = ROOT / "config/country_identity_registry.json"
DEFAULT_COUNTRIES = ROOT / "generated/town_country_asset_candidate/countries.json"
DEFAULT_ISO_REPORT = ROOT / "reports/iso-assignment-proposals.md"
DEFAULT_PERSIAN_REPORT = ROOT / "reports/persian-localization-proposals.md"
DEFAULT_FLAG_MAPPING = ROOT / "config/country_flag_mapping.json"
DEFAULT_OUTPUT = ROOT / "reports/country-data-inspection.xlsx"

ISO_HEADER = [
    "game_id", "Registry English name", "Candidate alpha-2",
    "Candidate alpha-3", "Reference English name", "Match type",
    "Authoritative evidence status", "Proposed disposition",
]
PERSIAN_HEADER = [
    "game_id", "Source ID", "English name", "Current Persian name",
    "Wikidata QID", "Candidate Persian name", "Candidate status", "Review status",
]


def read_json(path: Path, label: str) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"{label} file not found: {path}")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{label} must contain a JSON object: {path}")
    return value


def read_markdown_table(path: Path, header: list[str], label: str) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"{label} report not found: {path}")
    lines = path.read_text(encoding="utf-8").splitlines()
    expected = "| " + " | ".join(header) + " |"
    table_header = header
    try:
        start = lines.index(expected)
    except ValueError as exc:
        # The current proposal generator still emits the original seven-column
        # schema. Accept it without treating its Decision text as authoritative
        # verification; normalize both schemas to the workbook's internal shape.
        legacy_iso_header = [
            "game_id", "Registry English name", "Candidate alpha-2",
            "Candidate alpha-3", "Reference English name", "Match type", "Decision",
        ]
        legacy_expected = "| " + " | ".join(legacy_iso_header) + " |"
        if label != "ISO proposal" or legacy_expected not in lines:
            raise ValueError(f"{label} report has no expected table header: {path}") from exc
        table_header = legacy_iso_header
        start = lines.index(legacy_expected)
    rows: list[dict[str, str]] = []
    for line_number, line in enumerate(lines[start + 2 :], start=start + 3):
        if not line.startswith("|"):
            break
        cells = [cell.strip().replace("\\|", "|") for cell in line.strip().strip("|").split("|")]
        if len(cells) != len(table_header):
            raise ValueError(f"{label} report line {line_number}: expected {len(table_header)} columns, found {len(cells)}")
        row = dict(zip(table_header, cells, strict=True))
        if label == "ISO proposal" and "Decision" in row:
            row["Proposed disposition"] = row["Decision"]
            row["Authoritative evidence status"] = "Not verified; official ISO OBP review required"
        rows.append(row)
    if label == "ISO proposal":
        # Normalize both current and legacy schemas for workbook consumers.
        for row in rows:
            if "Proposed disposition" in row:
                row["Decision"] = row["Proposed disposition"]
    return rows


def text(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def write_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    *,
    table_name: str | None = None,
    freeze: str = "A2",
) -> None:
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([text(value) for value in row])
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for index, column in enumerate(ws.columns, start=1):
        sample = [str(cell.value or "") for cell in list(column)[:150]]
        max_length = max((len(value) for value in sample), default=10)
        ws.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 42)
    if table_name and rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def build_workbook(
    registry: dict[str, Any],
    countries_data: dict[str, Any] | None,
    iso_rows: list[dict[str, str]],
    persian_rows: list[dict[str, str]],
    flag_mapping: dict[str, Any],
    output: Path,
) -> dict[str, int]:
    records = registry.get("records")
    if not isinstance(records, list):
        raise ValueError("registry.records must be an array")
    country_list = countries_data.get("countries", []) if countries_data else []
    if not isinstance(country_list, list):
        raise ValueError("candidate countries.json field 'countries' must be an array")
    candidate_by_id = {
        str(row.get("game_id")): row for row in country_list
        if isinstance(row, dict) and row.get("game_id")
    }
    iso_by_id = {row["game_id"]: row for row in iso_rows if row.get("game_id")}
    persian_by_id = {row["game_id"]: row for row in persian_rows if row.get("game_id")}
    flag_records = flag_mapping.get("records", [])
    if not isinstance(flag_records, list):
        raise ValueError("flag mapping field 'records' must be an array")
    flags_by_id = {str(row.get("game_id")): row for row in flag_records if isinstance(row, dict) and row.get("game_id")}

    inventory_headers = [
        "game_id", "English name", "Current Persian name", "English aliases",
        "Persian aliases", "Current source ID", "Previous source IDs", "Geometry key",
        "Registry record status", "Entity type", "ISO alpha-2 current",
        "ISO alpha-3 current", "ISO assignment status", "Localization status",
        "ISO alpha-2 candidate", "ISO alpha-3 candidate", "ISO reference name",
        "ISO match type", "ISO proposal decision", "Persian candidate",
        "Persian candidate status", "Persian review status", "Wikidata QID",
        "Flag source code", "Flag mapping status", "Flag decision note",
        "Flag evidence", "Candidate compiled name", "Candidate compiled ISO3",
        "Candidate regions", "Candidate present",
    ]
    inventory_rows: list[list[Any]] = []
    iso_review_rows: list[list[Any]] = []
    entity_rows: list[list[Any]] = []
    persian_review_rows: list[list[Any]] = []
    flag_review_rows: list[list[Any]] = []
    quality_rows: list[list[Any]] = []
    source_attr_rows: list[dict[str, Any]] = []

    seen_ids: Counter[str] = Counter()
    for record in records:
        if not isinstance(record, dict):
            quality_rows.append(["ERROR", "", "Registry entry is not an object", "Inspect config/country_identity_registry.json"])
            continue
        gid = str(record.get("game_id") or "")
        seen_ids[gid] += 1
        source = record.get("source") if isinstance(record.get("source"), dict) else {}
        names = record.get("names") if isinstance(record.get("names"), dict) else {}
        iso = record.get("iso") if isinstance(record.get("iso"), dict) else {}
        iso_row = iso_by_id.get(gid, {})
        fa_row = persian_by_id.get(gid, {})
        flag = flags_by_id.get(gid, {})
        candidate = candidate_by_id.get(gid, {})
        attrs = candidate.get("source_attributes") if isinstance(candidate.get("source_attributes"), dict) else {}
        source_attr_rows.append({
            "game_id": gid,
            "English name": names.get("en"),
            "Current source ID": source.get("current_source_id"),
            **attrs,
        })
        inventory_rows.append([
            gid, names.get("en"), names.get("fa"), names.get("aliases_en", []),
            names.get("aliases_fa", []), source.get("current_source_id"),
            source.get("previous_source_ids", []), record.get("geometry_key"),
            record.get("status"), record.get("entity_type"), iso.get("alpha2"),
            iso.get("alpha3"), iso.get("assignment_status"),
            names.get("localization_status"), iso_row.get("Candidate alpha-2"),
            iso_row.get("Candidate alpha-3"), iso_row.get("Reference English name"),
            iso_row.get("Match type"), iso_row.get("Decision"),
            fa_row.get("Candidate Persian name"), fa_row.get("Candidate status"),
            fa_row.get("Review status"), fa_row.get("Wikidata QID"),
            flag.get("source_flag_code"), flag.get("status"), flag.get("decision_note"),
            flag.get("evidence", []), candidate.get("name"), candidate.get("iso3"),
            candidate.get("regions", []), "yes" if candidate else "no",
        ])
        iso_review_rows.append([
            gid, names.get("en"), iso.get("alpha2"), iso.get("alpha3"),
            iso.get("assignment_status"), iso_row.get("Candidate alpha-2"),
            iso_row.get("Candidate alpha-3"), iso_row.get("Reference English name"),
            iso_row.get("Match type"), iso_row.get("Decision"),
            "", "", "unreviewed",
        ])
        entity_rows.append([
            gid, names.get("en"), record.get("entity_type"), "", "", "unreviewed",
            "Record the classification decision and cite evidence. This worksheet does not adjudicate sovereignty or status.",
        ])
        persian_review_rows.append([
            gid, source.get("current_source_id"), names.get("en"), names.get("fa"),
            fa_row.get("Wikidata QID"), fa_row.get("Candidate Persian name"),
            fa_row.get("Candidate status"), fa_row.get("Review status"),
            "", "", "unreviewed",
        ])
        flag_review_rows.append([
            gid, names.get("en"), source.get("current_source_id"), record.get("entity_type"),
            iso.get("alpha2"), iso.get("alpha3"), flag.get("source_flag_code"),
            flag.get("status"), flag.get("decision_note"), flag.get("evidence", []),
            "", "", "unreviewed",
        ])
        reasons: list[str] = []
        if not gid:
            reasons.append("missing game_id")
        if record.get("status") != "approved":
            reasons.append(f"registry status={record.get('status')!r}")
        if record.get("entity_type") in (None, "", "unreviewed"):
            reasons.append("entity type needs review")
        if not iso.get("alpha2"):
            reasons.append("current ISO alpha-2 missing")
        if iso.get("assignment_status") != "verified":
            reasons.append(f"ISO status={iso.get('assignment_status')!r}")
        if names.get("localization_status") != "reviewed" or not names.get("fa"):
            reasons.append("Persian localization needs review")
        if gid not in iso_by_id:
            reasons.append("ISO proposal row missing")
        elif not iso_row.get("Candidate alpha-2") or iso_row.get("Candidate alpha-2") == "—":
            reasons.append("ISO alpha-2 candidate missing; inspect evidence")
        if gid not in persian_by_id:
            reasons.append("Persian proposal row missing")
        elif not fa_row.get("Candidate Persian name") or fa_row.get("Candidate Persian name") == "—":
            reasons.append("Persian candidate missing")
        if not flag.get("source_flag_code"):
            reasons.append("flag source code missing")
        if flag.get("status") != "reviewed":
            reasons.append(f"flag mapping status={flag.get('status')!r}")
        if countries_data is not None and not candidate:
            reasons.append("compiled candidate record missing")
        if record.get("geometry_key") != source.get("current_source_id"):
            reasons.append("geometry key differs from current source ID")
        if reasons:
            quality_rows.append([
                "REVIEW", gid, "; ".join(reasons),
                "Review the linked worksheet(s); do not auto-approve candidates.",
            ])

    registry_ids = {str(r.get("game_id")) for r in records if isinstance(r, dict) and r.get("game_id")}
    for gid in sorted(set(flags_by_id) - registry_ids):
        quality_rows.append(["ERROR", gid, "Flag mapping references unknown game_id", "Remove or correct only after checking the registry and mapping contract."])
    for gid in sorted(set(candidate_by_id) - registry_ids):
        quality_rows.append(["ERROR", gid, "Compiled candidate references unknown game_id", "Inspect compiler/registry crosswalk."])
    for gid, count in seen_ids.items():
        if count > 1:
            quality_rows.append(["ERROR", gid, f"Duplicate game_id in registry ({count} occurrences)", "Resolve through explicit identity review; never regenerate IDs."])
    for gid in sorted(registry_ids - set(flags_by_id)):
        quality_rows.append(["REVIEW", gid, "Flag mapping row missing", "Add an explicit mapping decision after human review."])
    for gid in sorted(registry_ids - set(candidate_by_id)) if countries_data is not None else []:
        quality_rows.append(["ERROR", gid, "No compiled candidate row for registry ID", "Inspect source crosswalk and compiler output."])

    # Use all attributes observed in the candidate dataset, without silently dropping fields.
    attr_headers = ["game_id", "English name", "Current source ID"] + sorted({
        key for row in source_attr_rows for key in row if key not in {"game_id", "English name", "Current source ID"}
    })
    attr_rows = [[row.get(key, "") for key in attr_headers] for row in source_attr_rows]

    wb = Workbook()
    intro = wb.active
    intro.title = "Read Me"
    intro.sheet_view.showGridLines = False
    intro["A1"] = "TownCountryAsset — Country Data Inspection Workbook"
    intro["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    intro["A1"].fill = PatternFill("solid", fgColor="244062")
    intro.merge_cells("A1:D1")
    intro["A3"] = "Generated (UTC)"
    intro["B3"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    intro["A4"] = "Registry records"
    intro["B4"] = len(records)
    intro["A5"] = "Compiled candidate records"
    intro["B5"] = len(country_list) if countries_data is not None else "not included"
    intro["A7"] = "Purpose"
    intro["B7"] = "Human inspection and review only. This workbook does not update project source files or approve any proposal."
    intro["A8"] = "Authoritative identity"
    intro["B8"] = "config/country_identity_registry.json; game_id is the persistent identity key."
    intro["A9"] = "Current source attributes"
    intro["B9"] = "The Source Attributes sheet contains all attributes present in candidate countries.json when that file is supplied."
    intro["A10"] = "Proposals"
    intro["B10"] = "ISO and Persian values are proposals. Treat names and matches as evidence, not proof or approval."
    intro["A11"] = "Flag policy"
    intro["B11"] = "A flag is a representation choice. Leave uncertain cases unreviewed; do not infer from names or unverified ISO codes."
    intro["A12"] = "Review columns"
    intro["B12"] = "Blank reviewer decision/evidence columns are intentionally left for humans to complete. Do not treat edits to this workbook as applied changes."
    intro["A13"] = "Data quality"
    intro["B13"] = "Review rows are prompts for inspection; some are expected while records remain in proposal/unreviewed state."
    intro.column_dimensions["A"].width = 26
    intro.column_dimensions["B"].width = 110
    for row in range(3, 14):
        intro.cell(row, 1).font = Font(bold=True)
        intro.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    intro.freeze_panes = "A3"

    write_sheet(wb, "Master Inventory", inventory_headers, inventory_rows, table_name="MasterInventory")
    write_sheet(
        wb, "ISO Review",
        ["game_id", "English name", "Current alpha-2", "Current alpha-3", "Current assignment status",
         "Candidate alpha-2", "Candidate alpha-3", "Reference name", "Match type", "Proposal decision",
         "Human decision", "Reviewer", "Review status"],
        iso_review_rows, table_name="ISOReview",
    )
    write_sheet(
        wb, "Entity Review",
        ["game_id", "English name", "Current entity type", "Proposed entity type",
         "Decision rationale", "Review status", "Guidance"],
        entity_rows, table_name="EntityReview",
    )
    write_sheet(
        wb, "Persian Review",
        ["game_id", "Source ID", "English name", "Current Persian name", "Wikidata QID",
         "Candidate Persian name", "Candidate status", "Source review status",
         "Human-approved Persian name", "Evidence / correction note", "Review status"],
        persian_review_rows, table_name="PersianReview",
    )
    write_sheet(
        wb, "Flag Review",
        ["game_id", "English name", "Source ID", "Entity type", "Current alpha-2", "Current alpha-3",
         "Proposed source flag code", "Mapping status", "Existing decision note", "Existing evidence",
         "Human decision / source code", "Decision rationale / evidence", "Review status"],
        flag_review_rows, table_name="FlagReview",
    )
    write_sheet(
        wb, "Data Quality",
        ["Severity", "game_id", "Finding", "Suggested inspection action"],
        quality_rows, table_name="DataQuality",
    )
    write_sheet(
        wb, "Source Attributes",
        attr_headers, attr_rows, table_name="SourceAttributes",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return {
        "registry_records": len(records),
        "candidate_records": len(country_list) if countries_data is not None else 0,
        "iso_proposal_rows": len(iso_rows),
        "persian_proposal_rows": len(persian_rows),
        "flag_mapping_rows": len(flag_records),
        "quality_findings": len(quality_rows),
        "source_attribute_columns": len(attr_headers),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export a multi-sheet, read-only country data inspection workbook.")
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY)
    parser.add_argument("--countries", type=Path, default=DEFAULT_COUNTRIES,
                        help="compiled candidate countries.json used to include compiled fields and all source attributes")
    parser.add_argument("--iso-report", type=Path, default=DEFAULT_ISO_REPORT)
    parser.add_argument("--persian-report", type=Path, default=DEFAULT_PERSIAN_REPORT)
    parser.add_argument("--flag-mapping", type=Path, default=DEFAULT_FLAG_MAPPING)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    countries = read_json(args.countries, "compiled candidate")
    registry = read_json(args.registry, "identity registry")
    flags = read_json(args.flag_mapping, "flag mapping")
    iso_rows = read_markdown_table(args.iso_report, ISO_HEADER, "ISO proposal")
    persian_rows = read_markdown_table(args.persian_report, PERSIAN_HEADER, "Persian localization")
    counts = build_workbook(registry, countries, iso_rows, persian_rows, flags, args.output)
    print(f"Inspection workbook written: {args.output}")
    for key, value in counts.items():
        print(f"{key}: {value}")
    print("Review-only output: no source records or approvals were changed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
