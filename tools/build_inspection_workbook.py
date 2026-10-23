#!/usr/bin/env python3
"""Build a read-only Excel inspection workbook for all country/entity asset data.

The workbook is a review artifact only. It never changes source data, mappings,
identity assignments, approved assets, or the game repository.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REGISTRY = Path("config/country_identity_registry.json")
DEFAULT_MAPPING = Path("config/country_flag_mapping.json")
DEFAULT_METADATA = Path("generated/town_country_asset_candidate/countries.json")
DEFAULT_ISO_REPORT = Path("reports/iso-assignment-proposals.md")
DEFAULT_PERSIAN_DIR = Path("raw_data/persian")
DEFAULT_OUTPUT = Path("reports/country-entity-inspection.xlsx")


def read_json(path: Path, required: bool = False) -> dict[str, Any]:
    if not path.is_file():
        if required:
            raise FileNotFoundError(path)
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"Expected a JSON object in {path}")
    return data


def latest_persian_candidates(directory: Path) -> dict[str, Any]:
    if not directory.is_dir():
        return {}
    candidates = sorted(directory.glob("wikidata-persian-candidates-*.json"))
    if not candidates:
        return {}
    return read_json(candidates[-1])


def parse_iso_report(path: Path) -> dict[str, dict[str, str]]:
    """Read the proposal table as evidence; proposals remain unapproved."""
    if not path.is_file():
        return {}
    result: dict[str, dict[str, str]] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip().replace("\\|", "|") for cell in line.strip().strip("|").split("|")]
        if len(cells) < 7 or cells[0] == "game_id" or set("".join(cells[0])) <= {"-", ":"}:
            continue
        if not cells[0].startswith("tc_"):
            continue
        result[cells[0]] = {
            "registry_english_name": cells[1],
            "candidate_alpha2": cells[2],
            "candidate_alpha3": cells[3],
            "reference_english_name": cells[4],
            "match_type": cells[5],
            "decision": cells[6],
        }
    return result


def json_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def set_sheet(ws: Any, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([json_cell(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25476A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in ws.columns:
        letter = column[0].column_letter
        sample_lengths = [len(str(cell.value or "")) for cell in column[:150]]
        width = min(max(max(sample_lengths, default=10) + 2, 12), 42)
        ws.column_dimensions[letter].width = width


def build(args: argparse.Namespace) -> dict[str, int]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Workbook generation requires openpyxl. Run: "
            "uv run --with openpyxl python tools/build_inspection_workbook.py"
        ) from exc

    registry = read_json(args.registry, required=True)
    mapping_doc = read_json(args.mapping, required=True)
    metadata_doc = read_json(args.metadata)
    records = registry.get("records")
    if not isinstance(records, list):
        raise ValueError("Registry records must be an array")
    mappings = mapping_doc.get("records", [])
    flags_by_id = {r.get("game_id"): r for r in mappings if isinstance(r, dict)}
    metadata_by_source = {
        str(c.get("id")): c for c in metadata_doc.get("countries", [])
        if isinstance(c, dict) and c.get("id") is not None
    }
    iso_by_id = parse_iso_report(args.iso_report)
    persian_doc = latest_persian_candidates(args.persian_dir)
    persian_by_id = {
        r.get("game_id"): r for r in persian_doc.get("rows", []) if isinstance(r, dict)
    }

    master: list[dict[str, Any]] = []
    identity: list[dict[str, Any]] = []
    iso_rows: list[dict[str, Any]] = []
    persian_rows: list[dict[str, Any]] = []
    flag_rows: list[dict[str, Any]] = []
    source_rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    source_field_names: set[str] = set()
    registry_ids: set[str] = set()

    for rec in records:
        if not isinstance(rec, dict):
            issues.append({"severity": "error", "game_id": "", "issue": "Registry record is not an object"})
            continue
        gid = str(rec.get("game_id") or "")
        source = rec.get("source") if isinstance(rec.get("source"), dict) else {}
        names = rec.get("names") if isinstance(rec.get("names"), dict) else {}
        iso = rec.get("iso") if isinstance(rec.get("iso"), dict) else {}
        source_id = str(source.get("current_source_id") or "")
        src = metadata_by_source.get(source_id, {})
        attrs = src.get("source_attributes") if isinstance(src.get("source_attributes"), dict) else {}
        flag = flags_by_id.get(gid, {})
        proposal = iso_by_id.get(gid, {})
        loc = persian_by_id.get(gid, {})
        if gid in registry_ids:
            issues.append({"severity": "error", "game_id": gid, "issue": "Duplicate game_id in identity registry"})
        registry_ids.add(gid)
        source_field_names.update(attrs.keys())

        flags_status = str(flag.get("status") or "mapping record missing")
        flag_code = flag.get("source_flag_code")
        loc_candidate = loc.get("candidate_persian_name", "")
        iso_status = str(iso.get("assignment_status") or "missing")
        entity_type = str(rec.get("entity_type") or "missing")
        localization_status = str(names.get("localization_status") or "missing")
        review_reasons = []
        if rec.get("status") != "approved":
            review_reasons.append("identity record not approved")
        if iso_status != "verified":
            review_reasons.append("ISO assignment not verified")
        if entity_type == "unreviewed":
            review_reasons.append("entity type unreviewed")
        if localization_status != "reviewed" or not names.get("fa"):
            review_reasons.append("Persian localization needs review")
        if flags_status != "reviewed" or not flag_code:
            review_reasons.append("flag mapping needs review")
        if not src:
            review_reasons.append("compiled source attributes unavailable; run pipeline first")

        master.append({
            "game_id": gid, "english_name": names.get("en", ""),
            "current_persian_name": names.get("fa", ""),
            "current_source_id": source_id, "previous_source_ids": source.get("previous_source_ids", []),
            "record_status": rec.get("status", ""), "entity_type": entity_type,
            "regions": rec.get("regions", []), "geometry_key": rec.get("geometry_key", ""),
            "registry_alpha2": iso.get("alpha2", ""), "registry_alpha3": iso.get("alpha3", ""),
            "iso_assignment_status": iso_status, "iso_candidate_alpha2": proposal.get("candidate_alpha2", ""),
            "iso_candidate_alpha3": proposal.get("candidate_alpha3", ""),
            "iso_match_type": proposal.get("match_type", ""), "iso_candidate_decision": proposal.get("decision", ""),
            "persian_candidate": loc_candidate, "persian_candidate_status": loc.get("candidate_status", ""),
            "localization_status": localization_status, "flag_source_code": flag_code or "",
            "flag_review_status": flags_status, "flag_decision_note": flag.get("decision_note", ""),
            "review_summary": "; ".join(review_reasons),
        })
        identity.append({
            "game_id": gid, "record_status": rec.get("status", ""), "english_name": names.get("en", ""),
            "aliases_en": names.get("aliases_en", []), "current_source_id": source_id,
            "previous_source_ids": source.get("previous_source_ids", []), "geometry_key": rec.get("geometry_key", ""),
            "regions": rec.get("regions", []), "registry_record_json": rec,
        })
        iso_rows.append({
            "game_id": gid, "english_name": names.get("en", ""),
            "registry_alpha2": iso.get("alpha2", ""), "registry_alpha3": iso.get("alpha3", ""),
            "assignment_status": iso_status, **proposal,
            "review_decision": "", "reviewer": "", "review_date": "", "review_evidence": "",
        })
        persian_rows.append({
            "game_id": gid, "english_name": names.get("en", ""),
            "current_persian_name": names.get("fa", ""), "candidate_persian_name": loc_candidate,
            "wikidata_qid": loc.get("wikidata_qid", ""), "candidate_status": loc.get("candidate_status", ""),
            "registry_localization_status": localization_status, "review_decision": "",
            "reviewer": "", "review_date": "", "review_evidence": "",
        })
        flag_rows.append({
            "game_id": gid, "english_name": names.get("en", ""), "entity_type": entity_type,
            "regions": rec.get("regions", []), "current_source_id": source_id,
            "source_flag_code": flag_code or "", "mapping_status": flags_status,
            "decision_note": flag.get("decision_note", ""), "evidence": flag.get("evidence", []),
            "review_decision": "", "reviewer": "", "review_date": "", "review_evidence": "",
        })
        source_rows.append({
            "game_id": gid, "english_name": names.get("en", ""), "current_source_id": source_id,
            "geometry_key": rec.get("geometry_key", ""), "source_geometry_type": src.get("source_geometry_type", ""),
            **attrs,
        })
        if not source_id:
            issues.append({"severity": "error", "game_id": gid, "issue": "Missing current_source_id"})
        if not names.get("en"):
            issues.append({"severity": "error", "game_id": gid, "issue": "Missing English name"})
        if not flag:
            issues.append({"severity": "warning", "game_id": gid, "issue": "No flag mapping record"})
        elif flags_status != "reviewed" or not flag_code:
            issues.append({"severity": "review", "game_id": gid, "issue": "Flag mapping unreviewed or source code empty"})
        if iso_status != "verified":
            issues.append({"severity": "review", "game_id": gid, "issue": "ISO assignment is not verified"})
        if entity_type == "unreviewed":
            issues.append({"severity": "review", "game_id": gid, "issue": "Entity type unreviewed"})
        if localization_status != "reviewed" or not names.get("fa"):
            issues.append({"severity": "review", "game_id": gid, "issue": "Persian localization missing or unreviewed"})
        if proposal and proposal.get("match_type") in {"explicit name variant", "ambiguous reference match"}:
            issues.append({"severity": "review", "game_id": gid, "issue": "ISO proposal uses name variant/ambiguous match; verify entity scope"})

    if len(registry_ids) != len(records):
        issues.append({"severity": "error", "game_id": "", "issue": "Registry contains duplicate/missing stable IDs"})
    flag_ids = set(flags_by_id)
    for gid in sorted(registry_ids - flag_ids):
        issues.append({"severity": "warning", "game_id": gid, "issue": "Registry ID missing from flag mapping file"})
    for gid in sorted(flag_ids - registry_ids):
        issues.append({"severity": "error", "game_id": gid, "issue": "Flag mapping refers to unknown game_id"})

    wb = Workbook()
    overview = wb.active
    overview.title = "Read Me"
    summary = [
        ["Country / Entity Data Inspection Workbook", ""],
        ["Generated at (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")],
        ["Registry records", len(records)],
        ["Unique game IDs", len(registry_ids)],
        ["Compiled source metadata available", bool(metadata_by_source)],
        ["ISO proposal rows loaded", len(iso_by_id)],
        ["Persian candidate rows loaded", len(persian_by_id)],
        ["Flag mapping rows", len(flags_by_id)],
        ["Data-quality/review findings", len(issues)],
        ["Purpose", "Human inspection and review only. This workbook does not approve or write back decisions."],
        ["Review safety", "Candidate values and existing values are shown in separate columns. Blank review-decision fields are intentional."],
        ["Recommended workflow", "Filter Master Inventory by review_summary; inspect the dedicated ISO, Entity & Identity, Persian Review, Flag Review, and Source Attributes sheets."],
        ["Source limitations", "Raw source attributes are populated from the latest candidate metadata file when available. If it is absent, run the data pipeline first; the workbook will identify missing source detail."],
        ["Output", str(args.output)],
    ]
    for row in summary:
        overview.append(row)
    overview.merge_cells("A1:B1")
    overview["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="25476A")
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 28
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 115
    for row in overview.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    overview.freeze_panes = "A2"

    set_sheet(wb.create_sheet("Master Inventory"), [
        "game_id", "english_name", "current_persian_name", "current_source_id", "previous_source_ids",
        "record_status", "entity_type", "regions", "geometry_key", "registry_alpha2", "registry_alpha3",
        "iso_assignment_status", "iso_candidate_alpha2", "iso_candidate_alpha3", "iso_match_type",
        "iso_candidate_decision", "persian_candidate", "persian_candidate_status", "localization_status",
        "flag_source_code", "flag_review_status", "flag_decision_note", "review_summary"
    ], master)
    set_sheet(wb.create_sheet("Entity & Identity"), [
        "game_id", "record_status", "english_name", "aliases_en", "current_source_id",
        "previous_source_ids", "geometry_key", "regions", "registry_record_json"
    ], identity)
    set_sheet(wb.create_sheet("ISO Review"), [
        "game_id", "english_name", "registry_alpha2", "registry_alpha3", "assignment_status",
        "registry_english_name", "candidate_alpha2", "candidate_alpha3", "reference_english_name",
        "match_type", "decision", "review_decision", "reviewer", "review_date", "review_evidence"
    ], iso_rows)
    set_sheet(wb.create_sheet("Persian Review"), [
        "game_id", "english_name", "current_persian_name", "candidate_persian_name", "wikidata_qid",
        "candidate_status", "registry_localization_status", "review_decision", "reviewer", "review_date",
        "review_evidence"
    ], persian_rows)
    set_sheet(wb.create_sheet("Flag Review"), [
        "game_id", "english_name", "entity_type", "regions", "current_source_id", "source_flag_code",
        "mapping_status", "decision_note", "evidence", "review_decision", "reviewer", "review_date",
        "review_evidence"
    ], flag_rows)
    source_headers = ["game_id", "english_name", "current_source_id", "geometry_key", "source_geometry_type"] + sorted(source_field_names)
    set_sheet(wb.create_sheet("Source Attributes"), source_headers, source_rows)
    set_sheet(wb.create_sheet("Data Quality"), ["severity", "game_id", "issue"], issues)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    return {
        "registry_records": len(records), "unique_game_ids": len(registry_ids),
        "source_attributes_rows": len(source_rows), "iso_proposal_rows": len(iso_by_id),
        "persian_candidate_rows": len(persian_by_id), "flag_mapping_rows": len(flags_by_id),
        "data_quality_findings": len(issues),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY)
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING)
    parser.add_argument("--metadata", type=Path, default=DEFAULT_METADATA,
                        help="Compiled candidate metadata JSON containing all source attributes")
    parser.add_argument("--iso-report", type=Path, default=DEFAULT_ISO_REPORT)
    parser.add_argument("--persian-dir", type=Path, default=DEFAULT_PERSIAN_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    args.registry = args.registry.expanduser()
    args.mapping = args.mapping.expanduser()
    args.metadata = args.metadata.expanduser()
    args.iso_report = args.iso_report.expanduser()
    args.persian_dir = args.persian_dir.expanduser()
    args.output = args.output.expanduser()
    try:
        counts = build(args)
    except (OSError, ValueError, RuntimeError, json.JSONDecodeError) as exc:
        parser.exit(2, f"error: {exc}\n")
    print(f"Inspection workbook written: {args.output}")
    for key, value in counts.items():
        print(f"{key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
